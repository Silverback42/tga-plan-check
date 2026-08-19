"""Unit-Tests fuer den Excel-Report-Export."""

from __future__ import annotations

from openpyxl import load_workbook

from app.models import DiffType
from app.reporting.excel import ReportRow, build_report

_SHEETS = ["Zusammenfassung", "Nur Schema", "Nur Grundriss", "Abweichungen"]


def _rows() -> list[ReportRow]:
    return [
        ReportRow(
            diff_type=DiffType.only_schema,
            severity="error",
            gewerk="HLK",
            label_schema="WRG 77",
            raum_schema="E.900",
            typ_schema="Lueftung",
        ),
        ReportRow(
            diff_type=DiffType.only_grundriss,
            severity="warning",
            gewerk="ELT",
            label_grundriss="KM 42",
            raum_grundriss="EG.500",
            typ_grundriss="Kaelte",
        ),
        ReportRow(
            diff_type=DiffType.attr_mismatch,
            severity="warning",
            gewerk="HLK",
            label_schema="RLT 01",
            label_grundriss="RLT 1",
            score=88.9,
            details={"schema_wert": "Lueftung", "grundriss_wert": "Pumpe"},
        ),
        ReportRow(
            diff_type=DiffType.room_mismatch,
            severity="warning",
            gewerk="HLK",
            label_schema="RLT 01",
            label_grundriss="RLT 1",
            score=88.9,
            details={"schema_wert": "E.801", "grundriss_wert": "EG.124"},
        ),
    ]


def test_build_report_erzeugt_alle_sheets(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    wb = load_workbook(target)
    assert wb.sheetnames == _SHEETS


def test_build_report_schreibt_datei_und_liefert_pfad(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "sub" / "r.xlsx")

    assert target.is_file()
    assert target.stat().st_size > 0


def test_zusammenfassung_enthaelt_counts_pro_typ(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Zusammenfassung"]
    values = {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}

    assert values["Nur im Schema"] == 1
    assert values["Nur im Grundriss"] == 1
    assert values["Attribut-Abweichung"] == 1
    assert values["Raum-Abweichung"] == 1
    assert values["Gesamt"] == 4


def test_zusammenfassung_zaehlt_pro_gewerk(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Zusammenfassung"]
    values = {row[0]: row[1] for row in sheet.iter_rows(values_only=True) if row[0]}

    assert values["HLK"] == 3
    assert values["ELT"] == 1


def test_projektname_steht_im_kopf(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Zusammenfassung"]
    assert sheet["A1"].value == "Projekt: Testprojekt"


def test_nur_schema_sheet_enthaelt_nur_passende_zeilen(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Nur Schema"]
    data = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(data) == 1
    assert data[0][1] == "WRG 77"
    assert data[0][2] == "E.900"


def test_nur_grundriss_sheet_nutzt_grundriss_spalten(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Nur Grundriss"]
    data = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(data) == 1
    assert data[0][1] == "KM 42"
    assert data[0][2] == "EG.500"


def test_abweichungen_sheet_enthaelt_beide_mismatch_typen(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    sheet = load_workbook(target)["Abweichungen"]
    data = list(sheet.iter_rows(min_row=2, values_only=True))

    assert len(data) == 2
    typen = {row[0] for row in data}
    assert typen == {"Attribut-Abweichung", "Raum-Abweichung"}
    # Werte-Spalten kommen aus details_json
    werte = {(row[4], row[5]) for row in data}
    assert ("Lueftung", "Pumpe") in werte
    assert ("E.801", "EG.124") in werte


def test_leerer_report_hat_sheets_aber_keine_datenzeilen(tmp_path):
    target = build_report([], "Leer", tmp_path / "r.xlsx")

    wb = load_workbook(target)
    assert wb.sheetnames == _SHEETS
    for name in _SHEETS[1:]:
        assert list(wb[name].iter_rows(min_row=2, values_only=True)) == []


def test_kopfzeile_ist_eingefroren(tmp_path):
    target = build_report(_rows(), "Testprojekt", tmp_path / "r.xlsx")

    wb = load_workbook(target)
    assert wb["Nur Schema"].freeze_panes == "A2"
