"""Excel-Export der Diff-Ergebnisse via openpyxl (Multi-Sheet)."""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.models import DiffType

# Kopfzeilen-Optik: einmal definiert, auf allen Sheets wiederverwendet
_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="4472C4")
_MAX_COL_WIDTH = 60

_DIFF_TYPE_LABELS = {
    DiffType.only_schema: "Nur im Schema",
    DiffType.only_grundriss: "Nur im Grundriss",
    DiffType.attr_mismatch: "Attribut-Abweichung",
    DiffType.room_mismatch: "Raum-Abweichung",
}


@dataclass(slots=True)
class ReportRow:
    """Eine denormalisierte Diff-Zeile fuer den Export."""

    diff_type: DiffType
    severity: str
    gewerk: str | None = None
    label_schema: str | None = None
    label_grundriss: str | None = None
    raum_schema: str | None = None
    raum_grundriss: str | None = None
    typ_schema: str | None = None
    typ_grundriss: str | None = None
    score: float | None = None
    details: dict = field(default_factory=dict)


def _write_header(sheet, headers: Sequence[str]) -> None:
    """Schreibt die Kopfzeile und friert sie ein."""
    sheet.append(list(headers))
    for cell in sheet[1]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    sheet.freeze_panes = "A2"


def _autosize(sheet) -> None:
    """Passt die Spaltenbreite an den laengsten Zellinhalt an."""
    for column in sheet.columns:
        longest = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        index = column[0].column
        sheet.column_dimensions[get_column_letter(index)].width = min(
            longest + 2, _MAX_COL_WIDTH
        )


def _add_summary_sheet(wb: Workbook, rows: Sequence[ReportRow], project_name: str) -> None:
    """Sheet 1: Counts pro Diff-Typ und pro Gewerk."""
    sheet = wb.active
    sheet.title = "Zusammenfassung"

    sheet.append([f"Projekt: {project_name}"])
    sheet["A1"].font = Font(bold=True, size=14)
    sheet.append([])

    sheet.append(["Diff-Typ", "Anzahl"])
    for cell in sheet[3]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for diff_type, label in _DIFF_TYPE_LABELS.items():
        count = sum(1 for r in rows if r.diff_type == diff_type)
        sheet.append([label, count])
    sheet.append(["Gesamt", len(rows)])
    sheet[sheet.max_row][0].font = Font(bold=True)
    sheet[sheet.max_row][1].font = Font(bold=True)

    sheet.append([])
    header_row = sheet.max_row + 1
    sheet.append(["Gewerk", "Anzahl"])
    for cell in sheet[header_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    # Gewerke stabil sortieren, damit der Report reproduzierbar bleibt
    for gewerk in sorted({r.gewerk for r in rows if r.gewerk}):
        sheet.append([gewerk, sum(1 for r in rows if r.gewerk == gewerk)])

    _autosize(sheet)


def _add_single_plan_sheet(
    wb: Workbook, rows: Iterable[ReportRow], title: str, diff_type: DiffType, label_col: str
) -> None:
    """Sheet 2/3: Anlagen, die nur in einem der beiden Plaene vorkommen."""
    sheet = wb.create_sheet(title)
    _write_header(sheet, ["Gewerk", label_col, "Raum", "Anlagentyp", "Schweregrad"])
    for row in rows:
        if row.diff_type != diff_type:
            continue
        if diff_type == DiffType.only_schema:
            label, raum, typ = row.label_schema, row.raum_schema, row.typ_schema
        else:
            label, raum, typ = row.label_grundriss, row.raum_grundriss, row.typ_grundriss
        sheet.append([row.gewerk, label, raum, typ, row.severity])
    _autosize(sheet)


def _add_mismatch_sheet(wb: Workbook, rows: Iterable[ReportRow]) -> None:
    """Sheet 4: Gematchte Paare mit Attribut- oder Raum-Abweichung."""
    sheet = wb.create_sheet("Abweichungen")
    _write_header(
        sheet,
        [
            "Diff-Typ",
            "Gewerk",
            "Label Schema",
            "Label Grundriss",
            "Wert Schema",
            "Wert Grundriss",
            "Score",
            "Schweregrad",
        ],
    )
    mismatch_types = (DiffType.attr_mismatch, DiffType.room_mismatch)
    for row in rows:
        if row.diff_type not in mismatch_types:
            continue
        sheet.append(
            [
                _DIFF_TYPE_LABELS[row.diff_type],
                row.gewerk,
                row.label_schema,
                row.label_grundriss,
                row.details.get("schema_wert"),
                row.details.get("grundriss_wert"),
                # Score auf 1 Nachkommastelle: im Report zaehlt Lesbarkeit
                None if row.score is None else round(row.score, 1),
                row.severity,
            ]
        )
    _autosize(sheet)


def build_report(rows: Sequence[ReportRow], project_name: str, target: Path) -> Path:
    """Schreibt die Diff-Zeilen als Multi-Sheet-Excel und liefert den Pfad."""
    wb = Workbook()
    _add_summary_sheet(wb, rows, project_name)
    _add_single_plan_sheet(
        wb, rows, "Nur Schema", DiffType.only_schema, "Label Schema"
    )
    _add_single_plan_sheet(
        wb, rows, "Nur Grundriss", DiffType.only_grundriss, "Label Grundriss"
    )
    _add_mismatch_sheet(wb, rows)

    target.parent.mkdir(parents=True, exist_ok=True)
    wb.save(target)
    return target
